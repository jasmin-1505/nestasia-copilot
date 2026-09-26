"""
Loads the 16 real rows from "Paid meta ads.xlsx" into PRODUCTION's
paid_ad_creative table only. Fixture is never touched by this script.

Reads directly from the spreadsheet (openpyxl) -- no hardcoded row values.
The real header row is row 4 (rows 1-3 are a title/blank rows Excel-side);
data starts at row 6.

Brand name mapping, decided by inspecting the actual sheet values against
the actual competitor_brand table, not assumed:
  - "Nestasia"                        -> brand (own brand)
  - "Home Centre"                     -> existing competitor_brand row
  - "Milton (incl. Treo, ProCook)"    -> existing competitor_brand "Milton"
                                          (same brand, sheet just used its
                                          fuller in-house name)
  - "Borosil"                         -> NEW competitor_brand (parked
                                          earlier, never created until now)
  - "IKEA India"                      -> NEW competitor_brand (new to this
                                          project entirely)
Creating Borosil/IKEA India here does NOT mean a full catalogue collector
exists for either -- platform is left NULL and a note makes this scope
explicit, so a future reader doesn't assume SKU-level coverage exists.

Every row's source_record uses source_type='ad_library_public' (this is
public data from Meta's own Ad Library, not synthetic) and collected_at =
the date this script actually ran (today), not any date from the sheet
itself -- the sheet's own "Start date" column is a property of the ad, not
of when this project observed it.

Known real-data oddity, preserved as-is rather than "fixed": row 11 (IKEA
India)'s "Landing page" cell contains a product name ("ASPEKT knife
sharpener, black - IKEA"), not a URL -- loaded verbatim into
landing_page_url, since guessing a real URL to replace it would be
inventing data that was never collected.

Usage:
    python db/load_paid_ad_creative.py
"""
import datetime
import os

import openpyxl
import psycopg2

ENV_FILE = os.path.join(os.path.dirname(__file__), "..", ".env")
XLSX_FILE = os.path.join(os.path.dirname(__file__), "..", "Paid meta ads.xlsx")

BRAND_NAME_MAP = {
    "Nestasia": ("brand", "Nestasia"),
    "Home Centre": ("competitor_brand", "Home Centre"),
    "Milton (incl. Treo, ProCook)": ("competitor_brand", "Milton"),
    "Borosil": ("competitor_brand", "Borosil"),
    "IKEA India": ("competitor_brand", "IKEA India"),
}


def _load_env():
    values = {}
    with open(ENV_FILE, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, _, val = line.partition("=")
            values[key.strip()] = val.strip()
    return values


def _connect(env, prefix):
    return psycopg2.connect(
        host=env[f"{prefix}_DB_HOST"], port=env[f"{prefix}_DB_PORT"],
        dbname=env[f"{prefix}_DB_NAME"], user=env[f"{prefix}_DB_USER"],
        password=env[f"{prefix}_DB_PASSWORD"], sslmode="require", connect_timeout=15,
    )


def _read_rows():
    wb = openpyxl.load_workbook(XLSX_FILE, data_only=True)
    ws = wb["Sheet1"]
    headers = [c.value for c in ws[4]]
    rows = []
    for raw in ws.iter_rows(min_row=6, values_only=True):
        record = dict(zip(headers, raw))
        if record.get("Brand") is None:
            continue
        rows.append(record)
    return rows


def main():
    env = _load_env()
    rows = _read_rows()
    print(f"Read {len(rows)} real rows from {XLSX_FILE}")

    unmapped = {r["Brand"] for r in rows} - set(BRAND_NAME_MAP)
    if unmapped:
        raise SystemExit(f"ABORTING -- unmapped brand name(s) in the sheet: {unmapped}")

    now = datetime.datetime.now(datetime.timezone.utc)

    conn = _connect(env, "PRODUCTION")
    conn.autocommit = False
    try:
        with conn.cursor() as cur:
            # ---- resolve/create brand + competitor_brand rows
            owner_ids = {}  # sheet brand name -> ("brand"|"competitor_brand", id)
            for sheet_name, (kind, canonical_name) in BRAND_NAME_MAP.items():
                if kind == "brand":
                    cur.execute("SELECT id FROM brand WHERE name = %s", (canonical_name,))
                    row = cur.fetchone()
                    if row is None:
                        raise SystemExit(f"ABORTING -- expected existing brand {canonical_name!r} not found")
                    owner_ids[sheet_name] = ("brand", row[0])
                else:
                    cur.execute("SELECT id FROM competitor_brand WHERE name = %s", (canonical_name,))
                    row = cur.fetchone()
                    if row is None:
                        note = ("Ad-creative data only (Meta Ad Library) -- no full catalogue "
                                "collector exists for this brand yet.")
                        cur.execute(
                            "INSERT INTO competitor_brand (name, platform, notes) VALUES (%s, NULL, %s) "
                            "RETURNING id",
                            (canonical_name, note)
                        )
                        new_id = cur.fetchone()[0]
                        print(f"  created NEW competitor_brand: {canonical_name!r} (id={new_id}), "
                              f"platform=NULL, notes={note!r}")
                        owner_ids[sheet_name] = ("competitor_brand", new_id)
                    else:
                        owner_ids[sheet_name] = ("competitor_brand", row[0])
                        print(f"  matched existing competitor_brand: {canonical_name!r} (id={row[0]})")

            # ---- load each row
            inserted = 0
            for r in rows:
                kind, owner_id = owner_ids[r["Brand"]]
                brand_id = owner_id if kind == "brand" else None
                competitor_brand_id = owner_id if kind == "competitor_brand" else None

                start_date = r["Start date"].date() if isinstance(r["Start date"], datetime.datetime) else r["Start date"]
                end_date = r["End date (if inactive)"]
                end_date = end_date.date() if isinstance(end_date, datetime.datetime) else end_date

                cur.execute(
                    "INSERT INTO source_record (source_type, source_url, collected_at, notes) "
                    "VALUES (%s, %s, %s, %s) RETURNING id",
                    ("ad_library_public", r["Ad library link / ID (source)"], now,
                     f"Sr no. {r['Sr no.']} in Paid meta ads.xlsx")
                )
                source_record_id = cur.fetchone()[0]

                cur.execute(
                    """
                    INSERT INTO paid_ad_creative (
                        brand_id, competitor_brand_id, platform, ad_library_url,
                        start_date, status, end_date, days_running, ad_format,
                        product_subcategory, hook_headline, offer_discount,
                        cta_button, landing_page_url, notes, theme, source_record_id
                    ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                    """,
                    (
                        brand_id, competitor_brand_id,
                        r["Platform"], r["Ad library link / ID (source)"],
                        start_date, r["Status"], end_date, r["Days running"], r["Format"],
                        r["Product subcategory"], r["Hook / headline (paraphrased)"],
                        r["Offer / discount"], r["CTA button"], r["Landing page"],
                        r["Notes"], r["Theme"], source_record_id,
                    )
                )
                inserted += 1

            print(f"Inserted {inserted} paid_ad_creative rows + {inserted} source_record rows.")

        conn.commit()
        print("COMMITTED.")
    except Exception:
        conn.rollback()
        print("ROLLED BACK -- no partial data left in production.")
        raise
    finally:
        conn.close()


if __name__ == "__main__":
    main()
